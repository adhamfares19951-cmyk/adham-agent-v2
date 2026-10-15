"""Word document builder for auto-generated job application packages."""
import io


def build_application_docx(title: str, cover_letter: str, resume_summary: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title or "Job Application Package", level=1)

    doc.add_heading("Cover Letter", level=2)
    for para in (cover_letter or "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_heading("Tailored CV Summary", level=2)
    for para in (resume_summary or "").split("\n"):
        if para.strip():
            style = "List Bullet" if para.strip().startswith(("-", "•")) else None
            doc.add_paragraph(para.strip().lstrip("-•").strip(), style=style)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
